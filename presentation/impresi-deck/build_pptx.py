#!/usr/bin/env python3
"""Generate a PowerPoint version of the FAP slide deck.

Mirrors the content of index.html (the Impresi deck) so the same presentation
is available as an editable, shareable .pptx file.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Dark theme palette (matching the Impresi deck)
BG      = RGBColor(0x0F, 0x11, 0x15)
BG_ALT  = RGBColor(0x16, 0x18, 0x1D)
BLUE    = RGBColor(0x7A, 0xA2, 0xF7)
TEXT    = RGBColor(0xC0, 0xCA, 0xF5)
GREEN   = RGBColor(0x9E, 0xCE, 0x6A)
RED     = RGBColor(0xF7, 0x76, 0x8E)
GOLD    = RGBColor(0xE0, 0xAF, 0x68)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
MUTED   = RGBColor(0xA9, 0xB1, 0xD6)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_slide(bg):
    slide = prs.slides.add_slide(BLANK)
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = bg
    return slide


def add_text(slide, text, *, x, y, w, h, size, color, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    f = run.font
    f.size = Pt(size)
    f.color.rgb = color
    f.bold = bold
    f.name = "Calibri"
    return box


def bullet(slide, items, *, x, y, w, size=22, gap=0.85):
    for i, (txt, color) in enumerate(items):
        add_text(slide, txt, x=x, y=y + i * gap, w=w, h=gap,
                 size=size, color=color)


# ---------- Slide 1: Title ----------
s = add_slide(BG)
add_text(s, "Functional Analytic Psychotherapy", x=0.9, y=2.4, w=11.5, h=1.4,
         size=44, color=BLUE, bold=True)
add_text(s, "Using the therapeutic relationship itself to create change",
         x=0.9, y=3.9, w=11.5, h=0.8, size=24, color=TEXT)
add_text(s, "Awareness  ·  Courage  ·  Love", x=0.9, y=4.9, w=11.5,
         h=0.7, size=20, color=GREEN)

# ---------- Slide 2: What is FAP ----------
s = add_slide(BG_ALT)
add_text(s, "What is FAP?", x=1.0, y=0.8, w=11, h=1, size=40, color=WHITE, bold=True)
add_text(s, "A behavioral psychotherapy developed by Robert Kohlenberg and "
            "Mavis Tsai (1991).", x=1.0, y=2.4, w=10.5, h=1.2, size=24, color=TEXT)
add_text(s, "Its core premise: meaningful change happens through a genuine, "
            "intense, and curative therapeutic relationship.",
         x=1.0, y=4.0, w=10.5, h=1.4, size=24, color=TEXT)

# ---------- Slide 3: CRBs ----------
s = add_slide(BG)
add_text(s, "Clinically Relevant Behaviors (CRBs)", x=0.8, y=0.7, w=12, h=1,
         size=34, color=BLUE, bold=True)
add_text(s, "The client's problems show up live, in the room, with the therapist.",
         x=0.8, y=2.0, w=11.5, h=0.8, size=22, color=MUTED)
bullet(s, [
    ("CRB1  —  problem behaviors as they occur in session", RED),
    ("CRB2  —  improvements that occur in session", GREEN),
    ("CRB3  —  the client's own talk about their behavior", GOLD),
], x=1.1, y=3.1, w=11, size=24, gap=1.1)

# ---------- Slide 4: The Five Rules ----------
s = add_slide(BG_ALT)
add_text(s, "The Five Rules", x=1.0, y=0.6, w=11, h=1, size=40, color=WHITE, bold=True)
bullet(s, [
    ("1  ·  Be aware — watch for CRBs", TEXT),
    ("2  ·  Be courageous — evoke CRBs", TEXT),
    ("3  ·  Be therapeutically loving — naturally reinforce CRB2s", TEXT),
    ("4  ·  Be aware of your impact on the client", TEXT),
    ("5  ·  Interpret functionally & generalize to daily life", TEXT),
], x=1.2, y=2.0, w=11, size=23, gap=0.95)

# ---------- Slide 5: Natural reinforcement ----------
s = add_slide(BG)
add_text(s, "Natural reinforcement", x=1.0, y=1.0, w=11, h=1, size=36,
         color=GREEN, bold=True)
add_text(s, "Change is driven by genuine, contingent responding to in-session "
            "improvements — not arbitrary praise. The therapist reacts as "
            "caring people would in the client's real life.",
         x=1.0, y=2.6, w=10.8, h=2.5, size=26, color=TEXT)

# ---------- Slide 6: ACL model ----------
s = add_slide(BG_ALT)
add_text(s, "The ACL model", x=1.0, y=0.8, w=11, h=1, size=40, color=BLUE, bold=True)
bullet(s, [
    ("Awareness — notice what matters, in yourself and the other", TEXT),
    ("Courage — risk emotional exposure and vulnerability", TEXT),
    ("Love — respond with genuine care and connection", TEXT),
], x=1.2, y=2.4, w=11, size=26, gap=1.3)

# ---------- Slide 7: Why it matters ----------
s = add_slide(BG)
add_text(s, "Why it matters", x=1.0, y=1.0, w=11, h=1, size=40, color=WHITE, bold=True)
add_text(s, "By working with behavior as it unfolds in the relationship, FAP "
            "builds new interpersonal repertoires that the client can carry "
            "into life outside the therapy room.",
         x=1.0, y=2.6, w=10.8, h=2.5, size=26, color=TEXT)

# ---------- Slide 8: Closing ----------
s = add_slide(BG)
add_text(s, "The relationship is the treatment.", x=0.8, y=2.9, w=11.7, h=1.3,
         size=40, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, "Kohlenberg & Tsai  ·  Functional Analytic Psychotherapy",
         x=0.8, y=4.3, w=11.7, h=0.7, size=20, color=GREEN, align=PP_ALIGN.CENTER)

out = "FAP-presentation.pptx"
prs.save(out)
print("Wrote", out, "with", len(prs.slides._sldIdLst), "slides")
